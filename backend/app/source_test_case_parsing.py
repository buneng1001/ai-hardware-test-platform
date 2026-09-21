import csv
import io
from pathlib import Path
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.source_test_case_models import SOURCE_CASE_FIELDS

MAX_IMPORT_BYTES = 2 * 1024 * 1024

FIELD_ALIASES = {
    "case_number": ("用例编号", "测试用例编号", "case_number", "case id", "id"),
    "title": ("测试用例标题", "用例标题", "title", "case title"),
    "priority": ("优先级", "priority"),
    "preconditions": ("预置条件", "前置条件", "preconditions"),
    "input": ("输入", "input"),
    "steps": ("操作步骤", "测试步骤", "steps"),
    "expected_result": ("预期结果", "expected result"),
    "test_type": ("测试类型", "test type"),
    "module": ("模块", "module"),
    "test_item": ("测试项", "test item"),
    "test_result": ("测试结果", "result"),
    "test_record": ("测试记录", "record"),
    "pre_test_notes": ("测试前备注信息", "测试前备注", "pre test notes"),
    "planned_execution_time": ("计划执行时间", "planned execution time"),
    "attachment": ("附件", "attachment"),
    "software_version": ("软件版本", "版本", "software version"),
}


def parse_source_case_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="导入文件不能超过 2 MiB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    raise HTTPException(status_code=422, detail="导入文件必须是 CSV 或 XLSX")


def infer_field_mapping(columns: list[str]) -> dict[str, str]:
    normalized_columns = {_normalize(column): column for column in columns}
    mapping: dict[str, str] = {}
    for field in SOURCE_CASE_FIELDS:
        for alias in FIELD_ALIASES[field]:
            matched = normalized_columns.get(_normalize(alias))
            if matched:
                mapping[field] = matched
                break
    return mapping


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=422, detail="CSV 必须使用 UTF-8 编码") from error
    reader = csv.DictReader(io.StringIO(text))
    columns = _validate_columns(reader.fieldnames)
    return columns, [{column: (row.get(column) or "").strip() for column in columns} for row in reader]


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise HTTPException(status_code=422, detail="Excel 文件无法读取") from error
    try:
        values = workbook.active.iter_rows(values_only=True)
        columns = _validate_columns([str(value).strip() if value is not None else "" for value in next(values, ())])
        return columns, [
            {column: str(value or "").strip() for column, value in zip(columns, row, strict=False)} for row in values
        ]
    finally:
        workbook.close()


def _validate_columns(columns: list[str] | None) -> list[str]:
    normalized = [column.strip() for column in columns or []]
    if not normalized or any(not column for column in normalized) or len(normalized) != len(set(normalized)):
        raise HTTPException(status_code=422, detail="表头不能为空且不能重复")
    return normalized


def _normalize(value: str) -> str:
    return "".join(value.lower().split()).replace("_", "").replace("-", "")
